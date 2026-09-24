from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

ROOT=Path('coffee_revision');FIG=ROOT/'figures'
main=ROOT/'Brazilian_coffee_manuscript_refined.docx';supp=ROOT/'Brazilian_coffee_supplementary_refined.docx'
names=['Figure1_adjustment_pathways.png','Figure2_geographic_reorganization.png','Figure3_climate_pressure.png','Figure4_environmental_redistribution.png','Figure5_water_productivity.png']
def replace_image(d,idx,name,width=6.3):
 sh=d.inline_shapes[idx];rid=sh._inline.graphic.graphicData.pic.blipFill.blip.embed
 d.part.related_parts[rid]._blob=(FIG/name).read_bytes()
 w,h=Image.open(FIG/name).size;sh.width=Inches(width);sh.height=Inches(width*h/w)
def change_caption(d,n,new):
 for p in d.paragraphs:
  if p.text.startswith('Figure '+str(n)+'.'):
   p.text=new;return
 raise ValueError(n)
d=Document(main)
existing_overview=any(p.text.startswith('Methodological overview.') for p in d.paragraphs)
offset=1 if existing_overview else 0
if existing_overview:replace_image(d,0,'FigureS1_analytical_workflow.png',width=6.25)
for i,name in enumerate(names):replace_image(d,i+offset,name,width=5.65 if i==1 else 6.25)
change_caption(d,1,'Figure 1. Contrasting production pathways for Brazilian Arabica and Canephora. (a–c) Annual national production, harvested area and aggregate productivity, 2012–2024. (d) Exact Shapley accounting contributions to production change between 2012–2015 and 2021–2024 (percentage points of baseline production). Bars are species coded; entry and exit refer to municipalities transitioning into or out of positive harvested area.')
change_caption(d,2,'Figure 2. Redistribution of Brazilian coffee harvested area. Each point is a municipality reference coordinate; symbol area in (a–d) scales with the municipality’s share of national species-specific harvested area in 2012–2015 or 2021–2024. Panels (e,f) show signed changes in municipal area share (blue, gain; orange, loss), with symbol area proportional to the absolute change. Arabica is shown in (a,b,e), Canephora in (c,d,f). Longitude and latitude are geographic coordinates; the maps do not locate individual fields.')
change_caption(d,3,'Figure 3. Climate exposure and seasonal rainfed water satisfaction. (a,b) Species-specific fixed-geography changes in mean temperature and annual days with Tmax ≥32 °C between 2012–2015 and 2021–2024. (c) Fixed-geography mean-temperature trend during 1981–2024, expressed per decade. (d) Changes in ETr/ETc (percentage points) by calendar window in the central standardized rainfed scenario; error ranges span 12 CAD × Kc scenarios and are parameter sensitivity envelopes, not confidence intervals. See Supplementary Table S6 for window definitions.')
change_caption(d,4,'Figure 4. Municipal environmental context of geographic redistribution. (a,b) Species-specific shares of harvested area by municipal reference-elevation band in 2012–2015 and 2021–2024. (c) Change in the area-weighted municipal reference elevation. (d) Redistribution-only shift in fixed 1981–2010 municipal mean temperature and hot-day climatology; hot-day values are divided by 20 for display and printed in original units. Municipal reference elevation is not field elevation; the independent coffee-pixel sensitivity is shown in Supplementary Figure S3.')
change_caption(d,5,'Figure 5. Water balance, productivity associations and chronological validation. (a) Central-scenario changes in ETr/ETc; whiskers span the 12 parameter scenarios and are not confidence intervals. (b) Conditional productivity associations for +100 mm rainfall (circle) or −0.10 water satisfaction (triangle), with exploratory 95% intervals. (c) Pooled 2020–2024 chronological squared-error improvement over a no-climate local-history comparator (positive is better). (d) Partial variance explained beyond local controls within sample. The estimates are neither causal damage measures nor operational forecast skill.')
for p in d.paragraphs:
 if p.text.startswith('Period differences in precipitation were spatially mixed (Fig. 3a–c; Supplementary Fig. S5).'):
  pass
 if 'Period differences in precipitation were spatially mixed (Fig. 3a–c; Supplementary Fig. S5).' in p.text:
  p.text=p.text.replace('Period differences in precipitation were spatially mixed (Fig. 3a–c; Supplementary Fig. S5).','Spatial differences in precipitation are documented in Supplementary Figure S5 (Fig. 3a,b shows the species-level fixed-geography contrasts).')
 if 'Fig. 2a–c' in p.text:p.text=p.text.replace('Fig. 2a–c','Fig. 2a,b,e')
 if 'Fig. 2d–f' in p.text:p.text=p.text.replace('Fig. 2d–f','Fig. 2c,d,f')
 if 'Municipal reference' in p.text and '(Fig. 4a–c)' in p.text:pass
if not existing_overview:
 heading=next(p for p in d.paragraphs if p.text=='2. Materials and methods')
 picture=d.add_paragraph();picture.alignment=WD_ALIGN_PARAGRAPH.CENTER
 picture.add_run().add_picture(str(FIG/'FigureS1_analytical_workflow.png'),width=Inches(6.25))
 heading._element.addnext(picture._element)
 lead=d.add_paragraph('Methodological overview. Municipal production, climate histories and an independent spatial sensitivity feed distinct accounting, exposure and productivity analyses. The detailed workflow is reproduced as Supplementary Figure S1.',style='Figure Caption')
 picture._element.addnext(lead._element)
d.save(main)

s=Document(supp);replace_image(s,0,'FigureS1_analytical_workflow.png',width=6.25)
for p in s.paragraphs:
 if p.text.startswith('Supplementary Figure S1.'):
  p.text='Supplementary Figure S1. Study design. IBGE species-resolved municipal production, NASA POWER daily climate and independent spatial sensitivity are joined at municipal reference points. The outputs quantify separate estimands: exact production-change accounting, fixed-geography exposure and redistribution, and conditional productivity associations with chronological validation. Arrows show analytical dependencies; they do not imply a causal climate effect.'
s.save(supp)

note=ROOT/'FIGURE_REVIEW.txt'
note.write_text('Figures 1–5 and Supplementary Figure S1 were rebuilt with a consistent DejaVu Sans font from the archived project CSV outputs and published supplementary Shapley accounting table.\n'
'Map panels show municipality reference coordinates rather than field locations; longitude and latitude axes are explicit. Fig. 3a–c fixed-geography summary values are transcribed from the manuscript and checked against its narrative.\n'
'The source annual climate CSV available through the archive was truncated in transit; no map was computed from the incomplete data. Supplementary map diagnostics remain in the document.\n'
'The 12-scenario ETr/ETc range represents parameter sensitivity, not a statistical confidence interval. Chart source: redraw_exact.py.\n')
with ZipFile(ROOT/'Brazilian_coffee_figures_review.zip','w',compression=ZIP_DEFLATED) as z:
 for file in sorted(FIG.iterdir()):
  if file.suffix in ('.png','.pdf'):z.write(file,arcname=file.name)
 z.write(ROOT/'redraw_exact.py',arcname='redraw_exact.py');z.write(note,arcname=note.name)
print('Embedded figures and workflow in both Word files.')
