"""Remove the workflow image from the main Word draft; keep Supplementary Fig. S1.

Usage: python move_workflow_to_supplement.py MAIN.docx SUPPLEMENT.docx
The second document is checked and left unchanged. The main document is edited in place.
"""
from pathlib import Path
import sys
from docx import Document


def main(main_path: Path, supplement_path: Path):
    supplement = Document(supplement_path)
    if not any(p.text.startswith("Supplementary Figure S1.") for p in supplement.paragraphs):
        raise RuntimeError("Supplementary Figure S1 caption is missing")
    main_doc = Document(main_path)
    heading = next(p for p in main_doc.paragraphs if p.text == "2. Materials and methods")
    successor = heading._element.getnext()
    if successor is not None and successor.xpath(".//w:drawing"):
        successor.getparent().remove(successor)
    for paragraph in list(main_doc.paragraphs):
        if paragraph.text.startswith("Methodological overview."):
            paragraph._element.getparent().remove(paragraph._element)
    first_methods = next(p for p in main_doc.paragraphs if p.text.startswith("Municipal coffee statistics were obtained"))
    if "Supplementary Fig. S1" not in first_methods.text:
        first_methods.add_run(" The analytical workflow is summarized in Supplementary Fig. S1.")
    main_doc.save(main_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(Path(sys.argv[1]), Path(sys.argv[2]))
